"""
Word 解析器

基于 python-docx。
- 段落：识别 Heading 样式（Heading 1~9 / 标题 N）与字体特征
- 表格：抽取为矩阵，交给 table_parser 做语义化
- 标题树：调用 heading_tree.build_heading_tree
"""
import re
from pathlib import Path
from typing import Optional

from loguru import logger

from .base import DocumentParser
from .models import Document, ParsedBlock, BlockType
from .heading_tree import build_heading_tree
from .table_parser import parse_table


_CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
           "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}


def _style_level(style_name: Optional[str]) -> Optional[int]:
    """从样式名提取标题级别。返回 None 表示非标题样式。"""
    if not style_name:
        return None

    m = re.match(r"Heading\s+(\d+)", style_name, re.IGNORECASE)
    if m:
        return int(m.group(1))

    m = re.match(r"标题\s*([0-9一二三四五六七八九十]+)", style_name)
    if m:
        token = m.group(1)
        return _CN_NUM.get(token, int(token) if token.isdigit() else 1)

    return None


class WordParser(DocumentParser):
    extensions = ("docx",)

    def parse(self, path: Path) -> Document:
        import docx

        doc = self._new_document(path, "word")
        document = docx.Document(str(path))

        # 1. 解析段落 → blocks
        blocks = self._parse_paragraphs(document)

        # 2. 构建标题树（含正文归属）
        doc.clean_blocks = blocks
        doc.heading_tree = build_heading_tree(blocks)
        doc.blocks = blocks

        # 3. 解析表格
        self._parse_tables(document, doc)

        logger.info(f"  Word 解析完成：{len(blocks)} 块，{len(doc.tables)} 张表")
        return doc

    def _parse_paragraphs(self, document) -> list:
        blocks = []
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_level = _style_level(para.style.name)

            # 字体特征（取第一个 run）
            font_size = None
            is_bold = False
            if para.runs:
                run = para.runs[0]
                if run.font.size is not None:
                    font_size = run.font.size.pt
                is_bold = bool(run.font.bold)

            block = ParsedBlock(
                block_type=BlockType.HEADING if style_level else BlockType.PARAGRAPH,
                text=text,
                font_size=font_size,
                is_bold=is_bold,
                metadata={
                    "style_level": style_level,
                    "style_name": para.style.name,
                },
            )
            blocks.append(block)

        return blocks

    def _parse_tables(self, document, doc: Document):
        for tbl in document.tables:
            matrix = []
            for row in tbl.rows:
                # python-docx 合并单元格会返回重复文本，此处为简化处理
                # （精确的 gridSpan/vMerge 展开见 TODO）
                matrix.append([cell.text.strip() for cell in row.cells])

            table = parse_table(
                raw_matrix=matrix,
                header_row_count=1,
                caption="",
                page=0,
            )
            doc.tables.append(table)
            doc.blocks.append(ParsedBlock(
                block_type=BlockType.TABLE,
                text=table.to_text(),
            ))
